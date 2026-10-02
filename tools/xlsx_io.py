"""work/text 의 JSON <-> translation/tms_message.xlsx (분류별 시트).

  python xlsx_io.py out     JSON -> xlsx   (ko 가 있으면 채워서 내보냄)
  python xlsx_io.py in      xlsx -> JSON   (ko 열만 반영)

열: file | i | label | ja | ko
"""
import os, sys
import text_io

ROOT = text_io.ROOT
XLSX = os.path.join(ROOT, 'translation', 'tms_message.xlsx')
COLS = ['file', 'i', 'label', 'ja', 'ko']


def category(p):
    return os.path.relpath(p, text_io.TEXT).split(os.sep)[0]


def cmd_out(args):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    n = 0
    for p in sorted(text_io.iter_json()):
        j = text_io.load_json(p)
        cat = category(p)
        ws = sheets.get(cat)
        if ws is None:
            ws = sheets[cat] = wb.create_sheet(cat[:31])
            ws.append(COLS)
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = 'A2'
            for col, w in zip('ABCDE', (34, 5, 26, 70, 70)):
                ws.column_dimensions[col].width = w
        short = j['file'][len(text_io.PREFIX):] if j['file'].startswith(text_io.PREFIX) else j['file']
        for en in j['entries']:
            ws.append([short, en['i'], en['label'], en['ja'], en.get('ko') or ''])
            for c in ws[ws.max_row][3:5]:
                c.alignment = Alignment(wrap_text=True, vertical='top')
            n += 1
    os.makedirs(os.path.dirname(XLSX), exist_ok=True)
    wb.save(XLSX)
    print('%d행 / %d시트 -> %s' % (n, len(sheets), XLSX))


def cmd_in(args):
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, read_only=True)
    ko_by = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        head = next(rows, None)
        if not head or list(head[:5]) != COLS:
            print('건너뜀(열 구성 다름):', ws.title)
            continue
        for r in rows:
            if not r or r[0] is None:
                continue
            ko = (r[4] or '').strip()
            if ko:
                ko_by[(str(r[0]), int(r[1]))] = ko
    n = changed = 0
    for p in text_io.iter_json():
        j = text_io.load_json(p)
        short = j['file'][len(text_io.PREFIX):] if j['file'].startswith(text_io.PREFIX) else j['file']
        dirty = False
        for en in j['entries']:
            ko = ko_by.get((short, en['i']), '')
            if ko and ko != (en.get('ko') or ''):
                en['ko'] = ko
                dirty = True
                changed += 1
            n += 1
        if dirty:
            text_io.save_json(p, j)
    print('xlsx 에서 %d개 번역 반영 (전체 %d행)' % (changed, n))


if __name__ == '__main__':
    {'out': cmd_out, 'in': cmd_in}[sys.argv[1]](sys.argv[2:])
